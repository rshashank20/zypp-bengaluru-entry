"""Build the outside-in Zypp Bengaluru fleet-sizing workbook.

This is a scenario model, not Zypp's internal forecast. Public facts and model
assumptions are kept separate; the blue cells are the assumptions to change.
Run from any working directory with: python3 build_sizing.py
"""
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "excel" / "zypp_sizing_model.xlsx"
OUT.parent.mkdir(parents=True, exist_ok=True)

# Styles
NAVY = "1F3864"
HDR_FILL = PatternFill("solid", fgColor=NAVY)
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
SEC_FILL = PatternFill("solid", fgColor="D9E2F3")
SEC_FONT = Font(bold=True, size=11)
INPUT_FILL = PatternFill("solid", fgColor="CCE5FF")
INPUT_FONT = Font(bold=True, size=11)
LBL_FONT = Font(size=10)
NOTE_FONT = Font(size=9, italic=True, color="595959")
TITLE_FONT = Font(bold=True, size=15, color=NAVY)
thin = Side(style="thin", color="B0B0B0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
COMMA = '#,##0'
PCT = '0.0%'
INR = '₹#,##0'
INR2 = '₹#,##0.00'


def setup(ws, title, widths):
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def header(ws, row, values):
    for i, value in enumerate(values, 1):
        c = ws.cell(row, i, value)
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def label(ws, row, value, col=1, bold=False):
    c = ws.cell(row, col, value)
    c.font = Font(size=10, bold=bold)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    c.border = BORDER


def input_cell(ws, row, value, fmt=None, col=2):
    c = ws.cell(row, col, value)
    c.font, c.fill, c.border = INPUT_FONT, INPUT_FILL, BORDER
    c.alignment = Alignment(horizontal="center", vertical="center")
    if fmt:
        c.number_format = fmt


def formula(ws, row, value, col=2, fmt=None, bold=False):
    c = ws.cell(row, col, value)
    c.font, c.border = Font(size=10, bold=bold), BORDER
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if fmt:
        c.number_format = fmt


def section(ws, row, title, ncols=4):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row, 1, title)
    c.font, c.fill, c.border = SEC_FONT, SEC_FILL, BORDER
    for col in range(2, ncols + 1):
        ws.cell(row, col).fill = SEC_FILL
        ws.cell(row, col).border = BORDER


def note(ws, row, text, ncols=4):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row, 1, text)
    c.font = NOTE_FONT
    c.alignment = Alignment(wrap_text=True, vertical="top")


wb = openpyxl.Workbook()
ws = wb.active
ws.title = "00_Cover"
setup(ws, "Zypp Electric — Bengaluru fleet-sizing stress test", [34, 76, 26, 25])
ws["A3"] = "Independent outside-in case study; not affiliated with Zypp Electric."
ws["A3"].font = NOTE_FONT
ws["A5"] = "Decision question"
ws["A5"].font = SEC_FONT
ws["A6"] = (
    "Zypp announced a 10,000-scooter Bengaluru target in May 2023. This workbook does not treat it as a current commitment. "
    "It tests what market share, active fleet and delivery productivity would be required if management revisited that scale."
)
ws["A6"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[6].height = 54
ws["A8"] = "How to use"
ws["A8"].font = SEC_FONT
for row, text in enumerate([
    "1. Change the blue cells on 01_Inputs; figures in this workbook are estimates or assumptions, not company forecasts.",
    "2. 02_Demand builds a city delivery-leg estimate from national quick-commerce orders and a reported Bengaluru food-delivery range.",
    "3. 03_Sizing translates TAM/SAM/SOM into physical scooters using daily productivity and availability assumptions.",
    "4. 04_Target tests the historic 10,000-scooter benchmark against the modeled market.",
    "5. 05_Scenarios varies the key assumptions. Its base case is linked to 01_Inputs; bear and bull inputs are illustrative.",
], 9):
    ws.cell(row, 1, text).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row, 1).font = LBL_FONT
    ws.row_dimensions[row].height = 31
note(ws, 16, "Important limitation: this is an outside-in sizing exercise. It cannot replace Zypp's zone-level demand, customer-contract, utilization, cost and vehicle-deployment data.")

# Inputs
ws = wb.create_sheet("01_Inputs")
setup(ws, "Inputs & assumptions — blue cells can be changed", [49, 22, 75, 35])
header(ws, 3, ["Input", "Base value", "Source / rationale", "Type / range"])
inputs = [
    ("National quick-commerce orders per day", 9_500_000, "Moneycontrol reported 9.5M national daily orders in Aug 2026.", "Public market estimate" , COMMA),
    ("Bengaluru share of national quick-commerce orders", 0.12, "Analyst assumption; not directly reported in the source set.", "Assumption: 10–15%", PCT),
    ("Bengaluru food-delivery orders per day", 550_000, "Moneycontrol reported an estimated 500K–600K Bengaluru food-delivery orders/day in Aug 2026; base uses midpoint.", "Public estimate range: 500K–600K", COMMA),
    ("Share of all delivery legs served through third-party fleet operators", 0.30, "Analyst assumption to define the serviceable slice; not a measured market statistic.", "Assumption: 20–45%", PCT),
    ("Zypp share of third-party legs", 0.217, "Calibrated midpoint: reported 6–7% Bengaluru last-mile share ÷ assumed 30% third-party share. Not independently observed.", "Calibrated assumption; test 12–28%", PCT),
    ("Deliveries per active rider/scooter per day", 26, "Model calibration assumption; aligns approximately with reported 5,000+ FY25 Bengaluru fleet and 6–7% market share under this market estimate.", "Assumption: 22–32", '0'),
    ("Share of physical fleet available on a typical day", 0.85, "Assumption for maintenance, downtime and other inactive vehicles.", "Assumption: 80–90%", PCT),
    ("Historic Bengaluru target used for the stress test (scooters)", 10_000, "Zypp publicly announced a target of 10,000 Bengaluru scooters in May 2023. This is historical, not a confirmed current target.", "Historical scenario benchmark", COMMA),
    ("Reported Bengaluru fleet floor in FY25 (vehicles)", 5_000, "Zypp reported more than 5,000 vehicles in Bengaluru in FY25; 5,000 is used as a floor for a rough consistency check.", "Public company-reported floor", COMMA),
]
for r, (name, value, rationale, typ, fmt) in enumerate(inputs, 4):
    label(ws, r, name)
    input_cell(ws, r, value, fmt)
    label(ws, r, rationale, 3)
    label(ws, r, typ, 4)
    ws.row_dimensions[r].height = 43
note(ws, 14, "Do not read calibrated inputs as independent evidence. The base-case Zypp share is anchored to company-reported market share and an assumed third-party share; this helps check scale, but does not independently validate the TAM.")

# Demand build
ws = wb.create_sheet("02_Demand")
setup(ws, "Market estimate — delivery legs per day", [57, 22, 38, 66])
header(ws, 3, ["Line", "Value", "Calculation", "Evidence / note"])
section(ws, 4, "A. Quick-commerce demand")
label(ws, 5, "National quick-commerce orders/day")
formula(ws, 5, "='01_Inputs'!B4", fmt=COMMA)
label(ws, 5, "Source input", 3)
label(ws, 5, "Moneycontrol, Aug 2026", 4)
label(ws, 6, "Bengaluru share")
formula(ws, 6, "='01_Inputs'!B5", fmt=PCT)
label(ws, 6, "Assumption", 3)
label(ws, 7, "Bengaluru quick-commerce orders/day", bold=True)
formula(ws, 7, "=B5*B6", fmt=COMMA, bold=True)
label(ws, 7, "National orders × city share", 3)
label(ws, 7, "The city-share input is not directly observed.", 4)
section(ws, 9, "B. Food-delivery demand")
label(ws, 10, "Bengaluru food-delivery orders/day")
formula(ws, 10, "='01_Inputs'!B6", fmt=COMMA)
label(ws, 10, "Public estimate; base uses midpoint", 3)
label(ws, 10, "Moneycontrol reported ~500K–600K orders/day in Bengaluru (Aug 2026).", 4)
section(ws, 12, "C. Total addressable demand")
label(ws, 13, "TAM — total modeled delivery legs/day", bold=True)
formula(ws, 13, "=B7+B10", fmt=COMMA, bold=True)
label(ws, 13, "Quick-commerce + food-delivery orders; one order ≈ one delivery leg", 3)
label(ws, 13, "This is a modeled proxy, not a published Bengaluru TAM.", 4)
note(ws, 15, "The total assumes one order corresponds to one delivery leg. It does not adjust for batching, multi-order trips, own-fleet riders, cancellations or differences in trip distance.")

# Sizing
ws = wb.create_sheet("03_Sizing")
setup(ws, "TAM / SAM / SOM and physical fleet requirement", [58, 22, 35, 58])
header(ws, 3, ["Line", "Value", "Calculation", "Interpretation"])
section(ws, 4, "A. Market layers")
label(ws, 5, "TAM — all modeled delivery legs/day")
formula(ws, 5, "='02_Demand'!B13", fmt=COMMA)
label(ws, 5, "Demand sheet", 3)
label(ws, 5, "All modeled two-wheeler delivery demand, before fleet-provider split.", 4)
label(ws, 6, "Third-party fleet share")
formula(ws, 6, "='01_Inputs'!B7", fmt=PCT)
label(ws, 6, "Assumption", 3)
label(ws, 7, "SAM — legs/day via third-party fleets", bold=True)
formula(ws, 7, "=B5*B6", fmt=COMMA, bold=True)
label(ws, 7, "TAM × third-party share", 3)
label(ws, 7, "Serviceable market proxy; the 30% input is an assumption.", 4)
label(ws, 8, "Zypp share of third-party fleet legs")
formula(ws, 8, "='01_Inputs'!B8", fmt=PCT)
label(ws, 8, "Calibrated assumption", 3)
label(ws, 9, "SOM — modeled Zypp legs/day", bold=True)
formula(ws, 9, "=B7*B8", fmt=COMMA, bold=True)
label(ws, 9, "SAM × Zypp share", 3)
label(ws, 9, "Base is calibrated to public share and assumed third-party penetration.", 4)
section(ws, 11, "B. Convert daily legs to physical scooters")
label(ws, 12, "Deliveries per active rider/scooter/day")
formula(ws, 12, "='01_Inputs'!B9", fmt='0')
label(ws, 12, "Assumption", 3)
label(ws, 13, "Physical fleet available on a typical day")
formula(ws, 13, "='01_Inputs'!B10", fmt=PCT)
label(ws, 13, "Assumption", 3)
label(ws, 14, "Scooters needed to serve modeled SOM", bold=True)
formula(ws, 14, "=ROUND(B9/B12/B13,0)", fmt=COMMA, bold=True)
label(ws, 14, "SOM ÷ deliveries/day ÷ availability", 3)
label(ws, 14, "Physical scooters required under the base assumptions.", 4)
label(ws, 15, "Scooters needed to serve the whole SAM")
formula(ws, 15, "=ROUND(B7/B12/B13,0)", fmt=COMMA)
label(ws, 15, "SAM ÷ deliveries/day ÷ availability", 3)
label(ws, 16, "Scooters needed to serve the whole TAM")
formula(ws, 16, "=ROUND(B5/B12/B13,0)", fmt=COMMA)
label(ws, 16, "TAM ÷ deliveries/day ÷ availability", 3)
note(ws, 18, "Fleet counts are capacity equivalents, not predictions of the number of vehicles Zypp will actually deploy. Actual trips per rider, operational uptime and contract demand are not public in sufficient detail.")

# Target test
ws = wb.create_sheet("04_Target")
setup(ws, "Stress-test of the historic 10,000-scooter benchmark", [60, 23, 39, 67])
header(ws, 3, ["Test", "Value", "Calculation", "Reading"])
label(ws, 4, "Historical target used for scenario (scooters)")
formula(ws, 4, "='01_Inputs'!B11", fmt=COMMA)
label(ws, 4, "May 2023 announcement", 3)
label(ws, 4, "Not a confirmed current Bengaluru target.", 4)
label(ws, 5, "Target-implied average legs/day")
formula(ws, 5, "=B4*'01_Inputs'!B9*'01_Inputs'!B10", fmt=COMMA, bold=True)
label(ws, 5, "Target × daily deliveries × availability", 3)
label(ws, 6, "Implied share of SAM")
formula(ws, 6, "=B5/'03_Sizing'!B7", fmt=PCT, bold=True)
label(ws, 6, "Target-implied legs ÷ SAM", 3)
label(ws, 6, "Market-capture requirement under the base case.", 4)
label(ws, 7, "Implied share of TAM")
formula(ws, 7, "=B5/'03_Sizing'!B5", fmt=PCT, bold=True)
label(ws, 7, "Target-implied legs ÷ TAM", 3)
label(ws, 8, "Difference vs fleet required for modeled SOM")
formula(ws, 8, "=B4-'03_Sizing'!B14", fmt=COMMA)
label(ws, 8, "Target fleet − modeled SOM fleet", 3)
label(ws, 9, "FY25 reported Bengaluru fleet floor")
formula(ws, 9, "='01_Inputs'!B12", fmt=COMMA)
label(ws, 9, "Company-reported 5,000+", 3)
label(ws, 10, "Modeled legs/day from FY25 fleet floor")
formula(ws, 10, "=B9*'01_Inputs'!B9*'01_Inputs'!B10", fmt=COMMA)
label(ws, 10, "Fleet floor × productivity × availability", 3)
label(ws, 11, "Implied share of TAM for FY25 fleet floor")
formula(ws, 11, "=B10/'03_Sizing'!B5", fmt=PCT)
label(ws, 11, "Modeled legs ÷ TAM", 3)
label(ws, 11, "Compare directionally with company-reported 6–7% market share; definitions may differ.", 4)
note(ws, 13, "The model's base case is calibrated in part to public market-share data. The FY25 share check is therefore a consistency check, not independent proof that the modeled TAM or fleet productivity is correct.")

# Scenarios
ws = wb.create_sheet("05_Scenarios")
setup(ws, "Bear / base / bull assumptions and outcomes", [55, 20, 20, 20, 54])
header(ws, 3, ["Driver", "Bear", "Base", "Bull", "Note"])
scenario_drivers = [
    ("Bengaluru share of national q-commerce", 0.10, "='01_Inputs'!B5", 0.15, PCT, "Assumption; national daily orders are anchored to Aug 2026 source."),
    ("Bengaluru food-delivery orders/day", 500_000, "='01_Inputs'!B6", 600_000, COMMA, "Public estimate range reported as 500K–600K/day."),
    ("Third-party fleet share of all delivery legs", 0.20, "='01_Inputs'!B7", 0.45, PCT, "Assumption; base value is not directly reported."),
    ("Zypp share of third-party legs", 0.12, "='01_Inputs'!B8", 0.28, PCT, "Base case calibrated from reported 6–7% city share and assumed 30% third-party share."),
    ("Deliveries per active rider/scooter/day", 22, "='01_Inputs'!B9", 32, '0', "Assumption; not a published operating KPI."),
    ("Physical fleet available on a typical day", 0.80, "='01_Inputs'!B10", 0.90, PCT, "Assumption for maintenance and downtime."),
]
for r, (name, bear, base, bull, fmt, note_text) in enumerate(scenario_drivers, 4):
    label(ws, r, name)
    input_cell(ws, r, bear, fmt)
    formula(ws, r, base, col=3, fmt=fmt)
    input_cell(ws, r, bull, fmt, col=4)
    label(ws, r, note_text, 5)
    ws.row_dimensions[r].height = 36
note(ws, 11, "Bear and bull settings are illustrative ranges, not company forecasts. Base cells link back to 01_Inputs so the model stays in sync.", ncols=5)
section(ws, 13, "Scenario outputs", ncols=5)
header(ws, 14, ["Output", "Bear", "Base", "Bull", "Calculation"])
out_rows = [
    ("Bengaluru q-commerce orders/day", COMMA, "National orders × city share"),
    ("TAM — delivery legs/day", COMMA, "Q-commerce + food-delivery orders"),
    ("SAM — third-party fleet legs/day", COMMA, "TAM × third-party share"),
    ("SOM — Zypp legs/day", COMMA, "SAM × Zypp share"),
    ("Physical fleet needed for SOM", COMMA, "SOM ÷ deliveries/day ÷ availability"),
    ("Target-implied share of SAM", PCT, "10,000 × deliveries/day × availability ÷ SAM"),
    ("Target-implied share of TAM", PCT, "10,000 × deliveries/day × availability ÷ TAM"),
]
for i, (name, fmt, calc) in enumerate(out_rows, 15):
    label(ws, i, name, bold=(i in (16, 17, 18, 19)))
    label(ws, i, calc, 5)
for col in ("B", "C", "D"):
    # Scenario driver rows: 4 city share, 5 food orders, 6 fleet share, 7 Zypp share,
    # 8 deliveries/active rider/day, 9 fleet availability.
    formula(ws, 15, f"='01_Inputs'!$B$4*{col}4", col=ord(col)-64, fmt=COMMA)
    formula(ws, 16, f"={col}15+{col}5", col=ord(col)-64, fmt=COMMA, bold=True)
    formula(ws, 17, f"={col}16*{col}6", col=ord(col)-64, fmt=COMMA, bold=True)
    formula(ws, 18, f"={col}17*{col}7", col=ord(col)-64, fmt=COMMA, bold=True)
    formula(ws, 19, f"=ROUND({col}18/{col}8/{col}9,0)", col=ord(col)-64, fmt=COMMA, bold=True)
    formula(ws, 20, f"=('01_Inputs'!$B$11*{col}8*{col}9)/{col}17", col=ord(col)-64, fmt=PCT, bold=True)
    formula(ws, 21, f"=('01_Inputs'!$B$11*{col}8*{col}9)/{col}16", col=ord(col)-64, fmt=PCT, bold=True)

# Sources
ws = wb.create_sheet("06_Sources")
setup(ws, "Source log — key external anchors", [9, 57, 48, 76])
header(ws, 3, ["ID", "Fact used", "Source and date", "URL"])
sources = [
    ("S1", "Historic Bengaluru 10,000-scooter announcement; 2,000 then deployed, 8,000 planned within two months.", "Indian Express / PTI, 2 May 2023", "https://indianexpress.com/article/business/startups/zypp-electric-e-scooters-bengaluru-next-2-months-8586983/lite/"),
    ("S2", "FY25: >5,000 Bengaluru vehicles; 6–7% Bengaluru last-mile delivery market share; national fleet about 20,000; 85–90% utilization claim.", "Business Standard, 24 Nov 2025", "https://www.business-standard.com/companies/news/zypp-electric-ipo-fleet-expansion-125112400958_1.html"),
    ("S3", "National quick-commerce platforms reached about 9.5 million orders/day.", "Moneycontrol, 17 Aug 2026", "https://www.moneycontrol.com/news/technology-startup/newsletters/MCTech3/tech3-quick-commerce-firms-cross-9m-daily-orders-canva-co-founder-cameron-adams-interview-and-more/"),
    ("S4", "Estimated 500,000–600,000 food-delivery orders/day in Bengaluru.", "Moneycontrol, 21 Aug 2026", "https://www.moneycontrol.com/news/business/startup/slowly-then-suddenly-rapidos-ownly-hits-50000-daily-food-delivery-orders-14012259.html"),
    ("S5", "FY25 revenue about ₹455 crore; 100M+ cumulative deliveries; 47% quick-commerce mix.", "Economic Times, 28 Apr 2025", "https://economictimes.indiatimes.com/tech/startups/zypp-electric-posts-rs-455-crore-revenue-in-fy25-up-50-yoy/articleshow/120697896.cms"),
    ("S6", "Reference Bziness XS battery-swappable scooter price of ₹57,750 under BaaS.", "Autocar Professional, 9 Dec 2025", "https://www.autocarpro.in/news/indofast-energy-and-quantum-energy-launch-affordable-electric-scooter-with-battery-swapping-130085"),
    ("S7", "EV logistics TCO benchmark: EVs can have lower TCO in high-utilization delivery fleets; this workbook's exact assumptions remain analyst inputs.", "ET Infra / Praxis Global Alliance, Jun 2025", "https://infra.economictimes.indiatimes.com/news/logistics/electric-vehicles-outperform-diesel-in-indian-logistics-cost-efficiency/121773106"),
]
for r, (sid, fact, source, url) in enumerate(sources, 4):
    label(ws, r, sid)
    label(ws, r, fact, 2)
    label(ws, r, source, 3)
    label(ws, r, url, 4)
    ws.row_dimensions[r].height = 43

for sheet in wb.worksheets:
    sheet.freeze_panes = "A4"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(OUT)
print(f"Wrote {OUT} ({OUT.stat().st_size:,} bytes)")
