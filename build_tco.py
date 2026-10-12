"""Build the per-scooter cost model: battery-swap EV reference vs petrol scooter.

This is a partial TCO scenario, not Zypp's actual negotiated operating economics.
Run from any working directory with: python3 build_tco.py
"""
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "excel" / "zypp_tco_model.xlsx"
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "1F3864"
HDR_FILL = PatternFill("solid", fgColor=NAVY)
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
SEC_FILL = PatternFill("solid", fgColor="D9E2F3")
SEC_FONT = Font(bold=True, size=11)
INPUT_FILL = PatternFill("solid", fgColor="CCE5FF")
INPUT_FONT = Font(bold=True, size=11)
WIN_FILL = PatternFill("solid", fgColor="C6EFCE")
WIN_FONT = Font(bold=True, size=11, color="006100")
LBL_FONT = Font(size=10)
NOTE_FONT = Font(size=9, italic=True, color="595959")
TITLE_FONT = Font(bold=True, size=15, color=NAVY)
thin = Side(style="thin", color="B0B0B0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
INR = '₹#,##0'
INR2 = '₹#,##0.00'
PCT = '0.0%'
COMMA = '#,##0'


def setup(ws, title, widths):
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def header(ws, row, values):
    for i, value in enumerate(values, 1):
        c = ws.cell(row, i, value)
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def label(ws, row, value, col=1, bold=False):
    c = ws.cell(row, col, value)
    c.font = Font(size=10, bold=bold)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    c.border = BORDER


def input_cell(ws, row, value, fmt=None, col=2):
    c = ws.cell(row, col, value)
    c.font, c.fill, c.border = INPUT_FONT, INPUT_FILL, BORDER
    c.alignment = Alignment(horizontal="center", vertical="center")
    if fmt:
        c.number_format = fmt


def formula(ws, row, value, col=2, fmt=None, bold=False, highlight=False):
    c = ws.cell(row, col, value)
    c.font = WIN_FONT if highlight else Font(size=10, bold=bold)
    if highlight:
        c.fill = WIN_FILL
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if fmt:
        c.number_format = fmt


def section(ws, row, title, ncols=4):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row, 1, title)
    c.font, c.fill, c.border = SEC_FONT, SEC_FILL, BORDER
    for col in range(2, ncols + 1):
        ws.cell(row, col).fill = SEC_FILL
        ws.cell(row, col).border = BORDER


def note(ws, row, text, ncols=4):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row, 1, text)
    c.font = NOTE_FONT
    c.alignment = Alignment(wrap_text=True, vertical="top")


wb = openpyxl.Workbook()
ws = wb.active
ws.title = "00_Cover"
setup(ws, "Zypp Electric — per-scooter cost comparison (partial TCO)", [32, 78, 25, 25])
ws["A3"] = "Independent outside-in model; not affiliated with Zypp Electric."
ws["A3"].font = NOTE_FONT
ws["A5"] = "Question"
ws["A5"].font = SEC_FONT
ws["A6"] = (
    "Using a publicly listed battery-swappable scooter as a reference, compare annualized vehicle, swap-plan, "
    "maintenance and insurance costs with a petrol scooter. This model does not include rider pay, financing, taxes, "
    "fleet overhead, downtime cost or resale value, and is not Zypp's actual unit economics."
)
ws["A6"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[6].height = 57
ws["A8"] = "How to use"
ws["A8"].font = SEC_FONT
for r, text in enumerate([
    "1. Change the blue cells on 01_Inputs. The EV base case assumes a ₹2,499/month plan with 75 kWh included.",
    "2. 02_EV_TCO flags if the average monthly energy requirement exceeds the selected plan allowance.",
    "3. 04_Comparison compares annual cost and calculates the break-even daily distance under the plan assumption.",
    "4. 05_Sensitivity covers 40–100 km per active operating day. Above the included energy allowance, a new plan price is needed.",
], 9):
    ws.cell(r, 1, text).font = LBL_FONT
    ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 31
note(ws, 14, "Reference vehicle price and plan price are public benchmarks, not verified Zypp purchase or fleet-contract terms. Subsidy is set to zero until the chosen model's eligibility and availability are confirmed.")

# Inputs
ws = wb.create_sheet("01_Inputs")
setup(ws, "Inputs & assumptions — blue cells can be changed", [47, 22, 79, 30])
header(ws, 3, ["Input", "Base value", "Source / rationale", "Type"])
inputs = [
    ("Reference battery-swap EV capex (ex-showroom)", 57750, "Bziness XS public BaaS reference price; not Zypp's actual procurement price.", "Public reference", INR),
    ("Petrol scooter capex (ex-showroom)", 92000, "Analyst assumption for a 110cc commuter scooter in Bengaluru.", "Assumption", INR),
    ("PM E-DRIVE subsidy on EV (if eligible)", 0, "Base set to zero; enter a subsidy only after confirming model-specific eligibility and availability.", "Assumption", INR),
    ("Vehicle life", 4, "Analyst assumption for commercial delivery use.", "Assumption (years)", '0'),
    ("Km per active operating day", 100, "Analyst assumption; compare under different daily distances.", "Assumption", '0'),
    ("Active operating days / year", 255, "300 potential operating days × 85% availability assumption.", "Assumption", '0'),
    ("Petrol price (Bengaluru)", 111, "Approximate public price around Oct 2026; actual daily price can vary.", "₹/L assumption", INR2),
    ("Petrol mileage", 45, "Analyst assumption for a 110cc scooter in urban delivery use.", "km/L assumption", '0'),
    ("EV consumption", 3.2, "Analyst assumption; test against the selected fleet vehicle under load.", "kWh/100km", '0.0'),
    ("BaaS subscription fee / month", 2499, "Published reference plan includes 75 kWh/month; actual fleet terms may differ.", "₹/month", INR),
    ("Energy included in BaaS plan / month", 75, "Published reference plan; average monthly usage must stay within this for the base calculation.", "kWh/month", '0'),
    ("EV maintenance", 0.30, "Analyst per-km allowance.", "₹/km", INR2),
    ("Petrol maintenance", 0.65, "Analyst per-km allowance.", "₹/km", INR2),
    ("EV insurance / year", 4500, "Analyst assumption.", "₹/year", INR),
    ("Petrol insurance / year", 3800, "Analyst assumption.", "₹/year", INR),
]
for r, (name, value, rationale, typ, fmt) in enumerate(inputs, 4):
    label(ws, r, name)
    input_cell(ws, r, value, fmt)
    label(ws, r, rationale, 3)
    label(ws, r, typ, 4)
    ws.row_dimensions[r].height = 40
note(ws, 20, "BaaS is a monthly plan, not a simple pay-per-kWh tariff. The model uses the plan fee if average monthly consumption stays within the included kWh. It does not estimate overage pricing; above the allowance, the model flags that more pricing data is required.")
I = "01_Inputs"

# EV cost build
ws = wb.create_sheet("02_EV_TCO")
setup(ws, "Reference BaaS EV — annual cost build", [48, 24, 43, 55])
header(ws, 3, ["Cost line", "Annual value", "Calculation", "Note"])
section(ws, 4, "Fixed costs")
label(ws, 5, "Vehicle amortization")
formula(ws, 5, f"=({I}!B4-{I}!B6)/{I}!B7", fmt=INR)
label(ws, 5, "(capex − confirmed subsidy) ÷ vehicle life", 3)
label(ws, 6, "Insurance")
formula(ws, 6, f"={I}!B17", fmt=INR)
label(ws, 6, "Annual insurance assumption", 3)
section(ws, 8, "Usage and energy")
label(ws, 9, "Annual km")
formula(ws, 9, f"={I}!B8*{I}!B9", fmt=COMMA)
label(ws, 9, "Km / active day × active operating days", 3)
label(ws, 10, "Annual energy use (kWh)")
formula(ws, 10, f"=B9*{I}!B12/100", fmt=COMMA)
label(ws, 10, "Annual km × kWh per 100 km ÷ 100", 3)
label(ws, 11, "Average energy use (kWh/month)")
formula(ws, 11, "=B10/12", fmt='0.0')
label(ws, 11, "Annual energy ÷ 12", 3)
label(ws, 12, "Annual BaaS subscription cost")
formula(ws, 12, f"={I}!B13*12", fmt=INR)
label(ws, 12, "Monthly fee × 12; applies only within included energy allowance", 3)
label(ws, 13, "Maintenance")
formula(ws, 13, f"=B9*{I}!B15", fmt=INR)
label(ws, 13, "Annual km × maintenance per km", 3)
section(ws, 15, "Total")
label(ws, 16, "Total annual EV cost", bold=True)
formula(ws, 16, f'=IF(B11<={I}!B14,B5+B6+B12+B13,"Check BaaS plan allowance")', fmt=INR, bold=True, highlight=True)
label(ws, 16, "Cost is only calculated if average monthly usage is within the selected plan's allowance.", 4)
label(ws, 17, "EV cost per km", bold=True)
formula(ws, 17, '=IF(ISNUMBER(B16),B16/B9,"")', fmt=INR2, bold=True)
label(ws, 17, "Annual EV cost ÷ annual km", 3)
note(ws, 19, "If monthly energy demand exceeds the plan allowance, this workbook deliberately does not invent an overage rate. Choose a higher-capacity plan or obtain an actual fleet tariff and rerun the model.")

# Petrol cost build
ws = wb.create_sheet("03_Petrol_TCO")
setup(ws, "Petrol scooter — annual cost build", [48, 24, 43, 40])
header(ws, 3, ["Cost line", "Annual value", "Calculation", "Note"])
section(ws, 4, "Fixed costs")
label(ws, 5, "Vehicle amortization")
formula(ws, 5, f"={I}!B5/{I}!B7", fmt=INR)
label(ws, 5, "Capex ÷ vehicle life", 3)
label(ws, 6, "Insurance")
formula(ws, 6, f"={I}!B18", fmt=INR)
label(ws, 6, "Annual insurance assumption", 3)
section(ws, 8, "Usage and fuel")
label(ws, 9, "Annual km")
formula(ws, 9, f"={I}!B8*{I}!B9", fmt=COMMA)
label(ws, 9, "Km / active day × active operating days", 3)
label(ws, 10, "Fuel cost")
formula(ws, 10, f"=B9/{I}!B11*{I}!B10", fmt=INR)
label(ws, 10, "Annual km ÷ km/L × petrol price", 3)
label(ws, 11, "Maintenance")
formula(ws, 11, f"=B9*{I}!B16", fmt=INR)
label(ws, 11, "Annual km × maintenance per km", 3)
section(ws, 13, "Total")
label(ws, 14, "Total annual petrol cost", bold=True)
formula(ws, 14, "=B5+B6+B10+B11", fmt=INR, bold=True, highlight=True)
label(ws, 15, "Petrol cost per km", bold=True)
formula(ws, 15, "=B14/B9", fmt=INR2, bold=True)
label(ws, 15, "Annual petrol cost ÷ annual km", 3)

# Comparison
ws = wb.create_sheet("04_Comparison")
setup(ws, "EV vs petrol — comparison under the base assumptions", [54, 26, 49, 45])
header(ws, 3, ["Metric", "Value", "Calculation", "Reading"])
label(ws, 4, "Annual cost — EV")
formula(ws, 4, "='02_EV_TCO'!B16", fmt=INR)
label(ws, 4, "Reference BaaS scooter", 3)
label(ws, 5, "Annual cost — petrol")
formula(ws, 5, "='03_Petrol_TCO'!B14", fmt=INR)
label(ws, 5, "Petrol reference scooter", 3)
label(ws, 6, "Annual cost difference (petrol − EV)", bold=True)
formula(ws, 6, '=IF(ISNUMBER(B4),B5-B4,"")', fmt=INR, bold=True, highlight=True)
label(ws, 6, "Positive means the EV costs less under these inputs", 3)
label(ws, 7, "Saving as % of petrol cost", bold=True)
formula(ws, 7, '=IF(ISNUMBER(B6),B6/B5,"")', fmt=PCT, bold=True, highlight=True)
label(ws, 8, "Cost per km — EV")
formula(ws, 8, "='02_EV_TCO'!B17", fmt=INR2)
label(ws, 9, "Cost per km — petrol")
formula(ws, 9, "='03_Petrol_TCO'!B15", fmt=INR2)
label(ws, 10, "Upfront vehicle-price difference (EV − petrol)")
formula(ws, 10, f"=({I}!B4-{I}!B6)-{I}!B5", fmt=INR)
label(ws, 10, "Negative means the reference EV has lower upfront capex", 3)
label(ws, 11, "Payback on an EV upfront premium")
formula(ws, 11, '=IF(B10<=0,"No premium to recover",IF(ISNUMBER(B6),B10/B6,"Needs energy-plan pricing"))')
section(ws, 13, "Break-even daily distance")
label(ws, 14, "EV variable cost per km within the plan")
formula(ws, 14, f"={I}!B15", fmt=INR2)
label(ws, 14, "Maintenance only; swap energy is in the monthly plan", 3)
label(ws, 15, "Petrol variable cost per km")
formula(ws, 15, f"={I}!B10/{I}!B11+{I}!B16", fmt=INR2)
label(ws, 16, "Annual fixed-cost gap (EV − petrol)")
formula(ws, 16, f"=(({I}!B4-{I}!B6)/{I}!B7+{I}!B17+{I}!B13*12)-({I}!B5/{I}!B7+{I}!B18)", fmt=INR)
label(ws, 16, "Includes vehicle amortization, insurance and the monthly swap plan", 3)
label(ws, 17, "Break-even km per active operating day", bold=True)
formula(ws, 17, f'=IF(B16<=0,0,IF(B15>B14,B16/(B15-B14)/{I}!B9,"No break-even in range"))', bold=True, highlight=True)
label(ws, 17, "Calculated only while monthly energy use stays within the plan", 3)
note(ws, 19, "This result changes with the scooter price, petrol price, mileage, usage, maintenance and BaaS plan. It is not a full company profitability comparison.")

# Sensitivity: base subscription includes up to 75 kWh/month
ws = wb.create_sheet("05_Sensitivity")
setup(ws, "Sensitivity — annual cost difference by daily distance", [28, 26, 20, 30, 48])
header(ws, 3, ["Km per active day", "Average kWh/month", "Annual saving ₹", "Winner", "Plan status"])
for r, km in enumerate([40, 60, 80, 100], 4):
    label(ws, r, km)
    formula(ws, r, f"={km}*'01_Inputs'!B9*'01_Inputs'!B12/100/12", col=2, fmt='0.0')
    petrol_total = f"('01_Inputs'!B5/'01_Inputs'!B7+'01_Inputs'!B18+({km}*'01_Inputs'!B9/'01_Inputs'!B11*'01_Inputs'!B10)+({km}*'01_Inputs'!B9*'01_Inputs'!B16))"
    ev_total = f"(('01_Inputs'!B4-'01_Inputs'!B6)/'01_Inputs'!B7+'01_Inputs'!B17+'01_Inputs'!B13*12+({km}*'01_Inputs'!B9*'01_Inputs'!B15))"
    formula(ws, r, f'=IF(B{r}<=\'01_Inputs\'!B14,{petrol_total}-{ev_total},"")', col=3, fmt=INR, bold=True, highlight=True)
    formula(ws, r, f'=IF(C{r}="","Needs plan pricing",IF(C{r}>0,"EV","Petrol"))', col=4)
    formula(ws, r, f'=IF(B{r}<=\'01_Inputs\'!B14,"Within included kWh","Plan cap exceeded")', col=5)
note(ws, 10, "Sensitivity stops at 100 km/day because this reference plan includes 75 kWh/month and the 100 km/day base uses about 68 kWh/month on average. Higher utilization needs a higher plan or verified overage price; do not extrapolate this table beyond the included energy.", ncols=5)

# Source log
ws = wb.create_sheet("06_Sources")
setup(ws, "Source log and assumptions", [9, 58, 44, 78])
header(ws, 3, ["ID", "Fact / input", "Source and date", "URL"])
sources = [
    ("T1", "Bziness XS battery-swappable scooter reference price: ₹57,750 ex-showroom under BaaS.", "Autocar Professional, 9 Dec 2025", "https://www.autocarpro.in/news/indofast-energy-and-quantum-energy-launch-affordable-electric-scooter-with-battery-swapping-130085"),
    ("T2", "BaaS plan: ₹999/month for 25 kWh or ₹2,499/month for 75 kWh.", "Quantum Energy public flyer", "https://quantumenergy.in/wp-content/uploads/2026/04/QE_EV_Bziness-XS_B2B-1.pdf"),
    ("T3", "Petrol price approx. ₹111/L; mileage and maintenance values are analyst assumptions.", "Public price trackers, Oct 2026; other metrics are assumptions", "https://www.business-standard.com/"),
    ("T4", "PM E-DRIVE scheme rules and extension; this model assumes no subsidy unless vehicle eligibility is verified.", "Ministry of Heavy Industries / PIB, Aug 2025–Oct 2026", "https://pmedrive.heavyindustries.gov.in/policy_document"),
]
for r, (sid, fact, source, url) in enumerate(sources, 4):
    label(ws, r, sid)
    label(ws, r, fact, 2)
    label(ws, r, source, 3)
    label(ws, r, url, 4)
    ws.row_dimensions[r].height = 42

for sheet in wb.worksheets:
    sheet.freeze_panes = "A4"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(OUT)
print(f"Wrote {OUT} ({OUT.stat().st_size:,} bytes)")
